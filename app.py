from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = Flask(__name__)
CORS(app)

EXCEL_FILE = "unsubscribe_log.xlsx"


@app.route("/unsubscribe", methods=["POST"])
def unsubscribe():
    data = request.get_json()

    email = data.get("email", "").strip()

    if not email:
        return jsonify({
            "success": False,
            "message": "Email is required"
        }), 400

    # Create Excel file if it does not exist
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Unsubscribe Log"

        sheet.append([
            "Email",
            "Status",
            "Source",
            "UnsubscribedAt"
        ])

        workbook.save(EXCEL_FILE)

    # Open existing Excel file
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Unsubscribe Log"]

    # Add unsubscribe record
    sheet.append([
        email,
        "UNSUBSCRIBED",
        "hashkrio-email",
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ])

    workbook.save(EXCEL_FILE)

    print("Unsubscribed:", email)

    return jsonify({
        "success": True,
        "message": "Unsubscribe request received"
    })


if __name__ == "__main__":
    app.run(debug=True, port=5000)